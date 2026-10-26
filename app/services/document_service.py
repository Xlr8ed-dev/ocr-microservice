import os
from typing import Any

from fastapi import APIRouter, HTTPException

from app.models import (
    DocumentOCRRequest,
    DocumentOCRResponse,
)

from app.services.ai_service import (
    AIServiceError,
)

from app.services.sso_service import (
    SSOServiceError,
    get_user_package,
)

from app.utils.file_utils import (
    delete_file,
    download_file,
)


router = APIRouter()


# =========================================================
# DOCUMENT EXTRACTION PROMPT
# =========================================================

DOCUMENT_PROMPT = """
You are a document processing assistant.

The OCR service has already extracted text from the document.

Your job is to process ONLY the extracted text provided to you.

Do not invent information.

Do not add information that is not present in the extracted text.

Maintain the meaning and structure of the original document.

Depending on the requested operation, perform the required
processing such as:

- read
- summarize
- extract
- analyze

Return a structured response suitable for the requested
operation.
"""


# =========================================================
# API
# =========================================================

@router.post(
    "",
    response_model=DocumentOCRResponse,
)
async def process_document(
    request: DocumentOCRRequest,
):
    """
    Generic document processing endpoint.

    Supported document types can include:

        PDF
        Image
        DOCX
        XLSX
        CSV
        TXT
        etc.

    Flow:

        Request
          ↓
        SSO
          ↓
        Package
          ↓
        Download document
          ↓
        Local document extraction
          ↓
        Extracted text
          ↓
        AI Service
          ↓
        Package-based model
          ↓
        Result
    """

    temp_file = None

    try:

        # -------------------------------------------------
        # Validate request
        # -------------------------------------------------

        if not request.file_url:
            raise HTTPException(
                status_code=400,
                detail="file_url is required",
            )

        # -------------------------------------------------
        # Get package from SSO
        # -------------------------------------------------

        try:

            package_details = await get_user_package(
                request.user_id
            )

        except SSOServiceError as exc:

            raise HTTPException(
                status_code=502,
                detail=str(exc),
            ) from exc

        # -------------------------------------------------
        # Download document
        # -------------------------------------------------

        try:

            temp_file = await download_file(
                request.file_url
            )

        except Exception as exc:

            raise HTTPException(
                status_code=400,
                detail=(
                    f"Unable to download document: "
                    f"{str(exc)}"
                ),
            ) from exc

        # -------------------------------------------------
        # Detect document type
        # -------------------------------------------------

        file_type = detect_file_type(
            temp_file
        )

        # -------------------------------------------------
        # Extract text
        # -------------------------------------------------

        try:

            extracted_text = await extract_document_text(
                file_path=temp_file,
                file_type=file_type,
            )

        except DocumentExtractionError as exc:

            raise HTTPException(
                status_code=422,
                detail=str(exc),
            ) from exc

        # -------------------------------------------------
        # Validate extracted text
        # -------------------------------------------------

        if not extracted_text.strip():

            raise HTTPException(
                status_code=422,
                detail=(
                    "No readable text could be extracted "
                    "from the document"
                ),
            )

        # -------------------------------------------------
        # User details
        # -------------------------------------------------

        user_details = {
            "user_id": request.user_id,
        }

        # -------------------------------------------------
        # Send extracted text to AI service
        # -------------------------------------------------

        ai_result = await process_document_with_ai(
            extracted_text=extracted_text,
            operation=request.operation,
            txn_id=request.txn_id,
            user_details=user_details,
            package_details=package_details,
        )

        # -------------------------------------------------
        # Get AI information
        # -------------------------------------------------

        ai_output = ai_result.get(
            "data",
            ai_result,
        )

        if not isinstance(ai_output, dict):
            ai_output = {
                "result": ai_output
            }

        # -------------------------------------------------
        # Return result
        # -------------------------------------------------

        return DocumentOCRResponse(
            success=True,
            txn_id=request.txn_id,
            file_name=os.path.basename(
                temp_file
            ),
            file_type=file_type,
            extracted_text=extracted_text,
            output=ai_output,
            package=_get_package_name(
                package_details
            ),
            ai_engine=ai_result.get(
                "model"
            ),
        )

    except HTTPException:
        raise

    except AIServiceError as exc:

        raise HTTPException(
            status_code=502,
            detail=str(exc),
        ) from exc

    except Exception as exc:

        raise HTTPException(
            status_code=500,
            detail=(
                f"Document processing failed: "
                f"{str(exc)}"
            ),
        ) from exc

    finally:

        # -------------------------------------------------
        # Remove downloaded temporary file
        # -------------------------------------------------

        if temp_file:
            delete_file(temp_file)


# =========================================================
# DOCUMENT TYPE
# =========================================================

def detect_file_type(
    file_path: str,
) -> str:
    """
    Detect document type from file extension.

    Later this can be replaced/enhanced with MIME detection
    or file signature detection.
    """

    extension = os.path.splitext(
        file_path
    )[1].lower()

    extension_map = {

        ".pdf": "pdf",

        ".jpg": "image",
        ".jpeg": "image",
        ".png": "image",
        ".webp": "image",
        ".bmp": "image",
        ".tiff": "image",
        ".tif": "image",

        ".docx": "docx",

        ".xlsx": "xlsx",
        ".xls": "xls",

        ".csv": "csv",

        ".txt": "txt",

        ".html": "html",
        ".htm": "html",
    }

    return extension_map.get(
        extension,
        "unknown",
    )


# =========================================================
# DOCUMENT TEXT EXTRACTION
# =========================================================

class DocumentExtractionError(Exception):
    """
    Raised when local document extraction fails.
    """


async def extract_document_text(
    *,
    file_path: str,
    file_type: str,
) -> str:

    """
    Extract text from a generic document.

    This function is intentionally kept independent from
    the AI service.

    The document extraction layer is responsible only for
    converting the source document into text.
    """

    if not os.path.exists(file_path):

        raise DocumentExtractionError(
            f"Document does not exist: {file_path}"
        )

    try:

        # -------------------------------------------------
        # PDF
        # -------------------------------------------------

        if file_type == "pdf":

            return extract_pdf_text(
                file_path
            )

        # -------------------------------------------------
        # Images
        # -------------------------------------------------

        if file_type == "image":

            return extract_image_text(
                file_path
            )

        # -------------------------------------------------
        # DOCX
        # -------------------------------------------------

        if file_type == "docx":

            return extract_docx_text(
                file_path
            )

        # -------------------------------------------------
        # XLSX
        # -------------------------------------------------

        if file_type == "xlsx":

            return extract_xlsx_text(
                file_path
            )

        # -------------------------------------------------
        # XLS
        # -------------------------------------------------

        if file_type == "xls":

            return extract_xls_text(
                file_path
            )

        # -------------------------------------------------
        # CSV
        # -------------------------------------------------

        if file_type == "csv":

            return extract_csv_text(
                file_path
            )

        # -------------------------------------------------
        # TXT
        # -------------------------------------------------

        if file_type == "txt":

            return extract_txt_text(
                file_path
            )

        # -------------------------------------------------
        # HTML
        # -------------------------------------------------

        if file_type == "html":

            return extract_html_text(
                file_path
            )

        raise DocumentExtractionError(
            f"Unsupported document type: {file_type}"
        )

    except DocumentExtractionError:
        raise

    except Exception as exc:

        raise DocumentExtractionError(
            f"Failed to extract document text: {str(exc)}"
        ) from exc


# =========================================================
# PDF
# =========================================================

def extract_pdf_text(
    file_path: str,
) -> str:

    try:

        import fitz

        document = fitz.open(
            file_path
        )

        pages = []

        try:

            for page in document:

                text = page.get_text()

                if text:
                    pages.append(text)

        finally:

            document.close()

        return "\n\n".join(
            pages
        ).strip()

    except ImportError as exc:

        raise DocumentExtractionError(
            "PyMuPDF is not installed"
        ) from exc


# =========================================================
# IMAGE
# =========================================================

def extract_image_text(
    file_path: str,
) -> str:

    """
    Generic image OCR.

    This is NOT the business-card OCR flow.

    Generic documents can use Tesseract or another local
    OCR engine.
    """

    try:

        from PIL import Image
        import pytesseract

        with Image.open(
            file_path
        ) as image:

            text = pytesseract.image_to_string(
                image,
                lang="eng",
            )

        return text.strip()

    except ImportError as exc:

        raise DocumentExtractionError(
            "Pillow or pytesseract is not installed"
        ) from exc

    except pytesseract.TesseractNotFoundError as exc:

        raise DocumentExtractionError(
            "Tesseract OCR engine was not found"
        ) from exc


# =========================================================
# DOCX
# =========================================================

def extract_docx_text(
    file_path: str,
) -> str:

    try:

        from docx import Document

        document = Document(
            file_path
        )

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        return "\n".join(
            paragraphs
        ).strip()

    except ImportError as exc:

        raise DocumentExtractionError(
            "python-docx is not installed"
        ) from exc


# =========================================================
# XLSX
# =========================================================

def extract_xlsx_text(
    file_path: str,
) -> str:

    try:

        from openpyxl import load_workbook

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        rows = []

        try:

            for worksheet in workbook.worksheets:

                rows.append(
                    f"[Sheet: {worksheet.title}]"
                )

                for row in worksheet.iter_rows(
                    values_only=True
                ):

                    values = [
                        str(value)
                        for value in row
                        if value is not None
                    ]

                    if values:
                        rows.append(
                            " | ".join(values)
                        )

        finally:

            workbook.close()

        return "\n".join(
            rows
        ).strip()

    except ImportError as exc:

        raise DocumentExtractionError(
            "openpyxl is not installed"
        ) from exc


# =========================================================
# XLS
# =========================================================

def extract_xls_text(
    file_path: str,
) -> str:

    try:

        import pandas as pd

        excel_file = pd.ExcelFile(
            file_path
        )

        rows = []

        for sheet_name in excel_file.sheet_names:

            rows.append(
                f"[Sheet: {sheet_name}]"
            )

            dataframe = pd.read_excel(
                file_path,
                sheet_name=sheet_name,
            )

            rows.append(
                dataframe.to_string(
                    index=False
                )
            )

        return "\n\n".join(
            rows
        ).strip()

    except ImportError as exc:

        raise DocumentExtractionError(
            "pandas/xlrd is not installed"
        ) from exc


# =========================================================
# CSV
# =========================================================

def extract_csv_text(
    file_path: str,
) -> str:

    try:

        import pandas as pd

        dataframe = pd.read_csv(
            file_path
        )

        return dataframe.to_string(
            index=False
        ).strip()

    except ImportError as exc:

        raise DocumentExtractionError(
            "pandas is not installed"
        ) from exc


# =========================================================
# TXT
# =========================================================

def extract_txt_text(
    file_path: str,
) -> str:

    try:

        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore",
        ) as file:

            return file.read().strip()

    except Exception as exc:

        raise DocumentExtractionError(
            f"Unable to read text file: {str(exc)}"
        ) from exc


# =========================================================
# HTML
# =========================================================

def extract_html_text(
    file_path: str,
) -> str:

    try:

        from bs4 import BeautifulSoup

        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore",
        ) as file:

            html = file.read()

        soup = BeautifulSoup(
            html,
            "html.parser",
        )

        return soup.get_text(
            separator="\n"
        ).strip()

    except ImportError as exc:

        raise DocumentExtractionError(
            "beautifulsoup4 is not installed"
        ) from exc


# =========================================================
# AI PROCESSING
# =========================================================

async def process_document_with_ai(
    *,
    extracted_text: str,
    operation: str,
    txn_id: str,
    user_details: dict[str, Any],
    package_details: dict[str, Any],
) -> dict[str, Any]:

    """
    Send extracted document text to AI service.

    IMPORTANT:

    This OCR service does not decide whether to use:

        local model
        OpenAI
        Gemini
        Ollama
        etc.

    AI service decides that using package_details.
    """

    prompt = f"""
{DOCUMENT_PROMPT}

Requested operation:

{operation}

Extracted document text:

-------------------------
{extracted_text}
-------------------------

Process the document according to the requested operation.
"""

    # -----------------------------------------------------
    # IMPORTANT
    # -----------------------------------------------------
    #
    # Your current ai_service.process_image() only accepts
    # an IMAGE.
    #
    # Generic document processing needs a text endpoint
    # in the AI service.
    #
    # Therefore this call represents the required future
    # AI-service method:
    #
    # process_text()
    #
    # -----------------------------------------------------

    from app.services.ai_service import process_text

    return await process_text(
        text=extracted_text,
        prompt=prompt,
        txn_id=txn_id,
        user_details=user_details,
        package_details=package_details,
    )


# =========================================================
# PACKAGE HELPER
# =========================================================

def _get_package_name(
    package_details: dict[str, Any],
) -> str | None:

    """
    Extract package name without assuming the exact SSO
    response structure.
    """

    return (
        package_details.get("package")
        or package_details.get("package_name")
        or package_details.get("name")
    )