"""Document text extraction.

Given a URL, download the file and return:
  - full text
  - per-page structure (Line objects with position / font hints)
  - page count
  - extraction method used

No AI. No DB. Pure I/O + CPU. Heading detection happens in the next
module (heading_detect.py) using the `lines` we return here.
"""

import io
import logging
import os
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf as fitz
import httpx
import pytesseract
from docx import Document as DocxDocument
from PIL import Image

from app.config import settings

logger = logging.getLogger(__name__)

# =========================================================
# TESSERACT CONFIG
# =========================================================

if settings.TESSERACT_CMD:
    pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD


# =========================================================
# ERRORS
# =========================================================

class DocumentExtractionError(Exception):
    """Raised when extraction fails for any reason."""


# =========================================================
# DATA SHAPES
# =========================================================

@dataclass
class Line:
    """One visual line of text, with position / font hints if available."""
    text: str
    page: int
    size: float | None = None
    bold: bool = False
    bbox: tuple[float, float, float, float] | None = None
    style: str | None = None


@dataclass
class ExtractedText:
    text: str  # full document text, joined by "\n"
    page_count: int
    lines: list[Line] = field(default_factory=list)
    extraction_method: str = "unknown"  # pdf_text | pdf_scan | docx | txt
    file_size_bytes: int = 0


# =========================================================
# CONSTANTS
# =========================================================

# If avg chars per page < this, treat PDF as scanned
SCANNED_PDF_CHARS_PER_PAGE = 100

# DPI for rendering scanned PDF pages to images
SCAN_RENDER_DPI = 200

# How long to wait for the download
DOWNLOAD_TIMEOUT = 120.0


# =========================================================
# PUBLIC ENTRY POINT
# =========================================================

async def extract_document(url: str) -> ExtractedText:
    """Download the file from `url` and extract text + structure.

    Detects the file type by extension and, for PDFs, by inspecting the
    text layer. Falls back to Tesseract OCR for scanned PDFs.

    Raises DocumentExtractionError on any failure.
    """
    if not url:
        raise DocumentExtractionError("Empty source URL")

    tmp_path: str | None = None
    try:
        tmp_path, file_size = await _download_to_temp(url)
        logger.info("Downloaded %s bytes from %s", file_size, url)

        ext = _detect_extension(url, tmp_path)
        logger.info("Detected file type: %s", ext)

        if ext == "pdf":
            result = _extract_pdf(tmp_path)
        elif ext == "docx":
            result = _extract_docx(tmp_path)
        elif ext == "doc":
            result = _extract_doc(tmp_path)
        elif ext == "txt":
            result = _extract_txt(tmp_path)
        else:
            raise DocumentExtractionError(f"Unsupported file type: {ext}")

        result.file_size_bytes = file_size
        return result

    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except OSError:
                pass


# =========================================================
# DOWNLOAD
# =========================================================

async def _download_to_temp(url: str) -> tuple[str, int]:
    """Stream the file to a temp file. Returns (path, size_bytes)."""
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".bin")
    tmp_path = tmp.name
    tmp.close()

    total = 0
    try:
        async with httpx.AsyncClient(timeout=DOWNLOAD_TIMEOUT) as client:
            async with client.stream("GET", url) as response:
                response.raise_for_status()
                with open(tmp_path, "wb") as f:
                    async for chunk in response.aiter_bytes(chunk_size=64 * 1024):
                        f.write(chunk)
                        total += len(chunk)
    except httpx.HTTPError as exc:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass
        raise DocumentExtractionError(f"Download failed: {exc}") from exc

    return tmp_path, total


def _detect_extension(url: str, local_path: str) -> str:
    """Prefer the URL extension; fall back to magic-byte sniffing."""
    url_path = url.split("?", 1)[0].lower()
    for ext in ("pdf", "docx", "doc", "txt"):
        if url_path.endswith("." + ext):
            return ext

    with open(local_path, "rb") as f:
        head = f.read(8)

    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        # ZIP-based — could be docx, xlsx, pptx. Assume docx.
        return "docx"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        # OLE2 compound file — legacy .doc
        return "doc"

    return "txt"


# =========================================================
# PDF
# =========================================================

def _extract_pdf(path: str) -> ExtractedText:
    """Extract a PDF. Auto-detects text vs scanned."""
    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise DocumentExtractionError(f"Cannot open PDF: {exc}") from exc

    try:
        page_count = doc.page_count
        if page_count == 0:
            raise DocumentExtractionError("PDF has no pages")

        lines, page_texts = _pdf_extract_lines(doc)
        total_chars = sum(len(t) for t in page_texts)
        avg_chars = total_chars / page_count if page_count else 0

        if avg_chars >= SCANNED_PDF_CHARS_PER_PAGE:
            logger.info("PDF treated as text (%.1f chars/page avg)", avg_chars)
            return ExtractedText(
                text="\n".join(page_texts),
                page_count=page_count,
                lines=lines,
                extraction_method="pdf_text",
            )

        logger.info(
            "PDF looks scanned (%.1f chars/page avg) — OCR'ing with Tesseract",
            avg_chars,
        )
        return _pdf_extract_via_ocr(doc)

    finally:
        doc.close()


def _pdf_extract_lines(doc: "fitz.Document") -> tuple[list[Line], list[str]]:
    """Extract per-line text + font hints from a text-based PDF."""
    all_lines: list[Line] = []
    page_texts: list[str] = []

    for page_num, page in enumerate(doc, start=1):
        try:
            blocks = page.get_text("dict")["blocks"]
        except Exception:
            page_texts.append("")
            continue

        page_line_buffer: list[str] = []

        for block in blocks:
            if block.get("type") != 0:  # 0 = text block
                continue
            for line in block.get("lines", []):
                spans = line.get("spans", [])
                if not spans:
                    continue

                line_text = "".join(s.get("text", "") for s in spans).strip()
                if not line_text:
                    continue

                total_len = 0
                weighted_size = 0.0
                is_bold = False
                for span in spans:
                    span_text = span.get("text", "")
                    slen = len(span_text)
                    total_len += slen
                    weighted_size += span.get("size", 0.0) * slen
                    if span.get("flags", 0) & (1 << 4):  # bit 4 = bold
                        is_bold = True

                avg_size = weighted_size / total_len if total_len else None

                bbox_raw = line.get("bbox")
                bbox = (
                    (bbox_raw[0], bbox_raw[1], bbox_raw[2], bbox_raw[3])
                    if bbox_raw else None
                )

                all_lines.append(Line(
                    text=line_text,
                    page=page_num,
                    size=avg_size,
                    bold=is_bold,
                    bbox=bbox,
                ))
                page_line_buffer.append(line_text)

        page_texts.append("\n".join(page_line_buffer))

    return all_lines, page_texts


def _pdf_extract_via_ocr(doc: "fitz.Document") -> ExtractedText:
    """Render each page to an image and OCR with Tesseract."""
    all_lines: list[Line] = []
    page_texts: list[str] = []

    for page_num, page in enumerate(doc, start=1):
        try:
            pix = page.get_pixmap(dpi=SCAN_RENDER_DPI)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
        except Exception as exc:
            logger.warning("Cannot render page %d: %s", page_num, exc)
            page_texts.append("")
            continue

        try:
            data = pytesseract.image_to_data(
                img,
                output_type=pytesseract.Output.DICT,
            )
        except Exception as exc:
            raise DocumentExtractionError(
                f"Tesseract failed on page {page_num}: {exc}"
            ) from exc

        line_groups: dict[tuple[int, int, int], list[dict]] = {}
        for i, word in enumerate(data["text"]):
            word = (word or "").strip()
            if not word:
                continue
            key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
            line_groups.setdefault(key, []).append({
                "text": word,
                "left": data["left"][i],
                "top": data["top"][i],
                "width": data["width"][i],
                "height": data["height"][i],
            })

        page_lines: list[str] = []
        for key in sorted(line_groups.keys()):
            words = line_groups[key]
            text = " ".join(w["text"] for w in words).strip()
            if not text:
                continue

            left = min(w["left"] for w in words)
            top = min(w["top"] for w in words)
            right = max(w["left"] + w["width"] for w in words)
            bottom = max(w["top"] + w["height"] for w in words)
            height = bottom - top

            all_lines.append(Line(
                text=text,
                page=page_num,
                size=float(height),
                bold=False,
                bbox=(float(left), float(top), float(right), float(bottom)),
            ))
            page_lines.append(text)

        page_texts.append("\n".join(page_lines))

    return ExtractedText(
        text="\n".join(page_texts),
        page_count=doc.page_count,
        lines=all_lines,
        extraction_method="pdf_scan",
    )


# =========================================================
# DOCX
# =========================================================

def _extract_docx(path: str) -> ExtractedText:
    """Extract text + style hints from a .docx file."""
    try:
        doc = DocxDocument(path)
    except Exception as exc:
        raise DocumentExtractionError(f"Cannot open DOCX: {exc}") from exc

    lines: list[Line] = []
    page_texts: list[str] = []
    page_num = 1

    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue

        style_name = None
        try:
            style_name = para.style.name if para.style else None
        except Exception:
            style_name = None

        is_heading = bool(style_name and style_name.lower().startswith("heading"))
        bold = is_heading or _paragraph_is_bold(para)

        lines.append(Line(
            text=text,
            page=page_num,
            size=None,
            bold=bold,
            bbox=None,
            style=style_name,
        ))

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            row_text = " | ".join(cells)
            if row_text.strip(" |"):
                lines.append(Line(text=row_text, page=page_num))

    page_texts.append("\n".join(l.text for l in lines))

    return ExtractedText(
        text="\n".join(page_texts),
        page_count=1,
        lines=lines,
        extraction_method="docx",
    )


def _paragraph_is_bold(para) -> bool:
    try:
        for run in para.runs:
            if run.bold:
                return True
    except Exception:
        pass
    return False


# =========================================================
# DOC (legacy) — deferred to v2
# =========================================================

def _extract_doc(path: str) -> ExtractedText:
    """Legacy .doc support is deferred to v2."""
    raise DocumentExtractionError(
        "Legacy .doc files are not supported yet. "
        "Please save the document as .docx or .pdf and re-upload."
    )


# =========================================================
# TXT
# =========================================================

def _extract_txt(path: str) -> ExtractedText:
    """Read a plain text file. Tries UTF-8 (with BOM handling), falls back to latin-1."""
    try:
        text = Path(path).read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        text = Path(path).read_text(encoding="latin-1")
    except Exception as exc:
        raise DocumentExtractionError(f"Cannot read TXT: {exc}") from exc

    return _flat_text_result(text, "txt")


# =========================================================
# SHARED HELPERS
# =========================================================

def _flat_text_result(text: str, method: str) -> ExtractedText:
    """Build an ExtractedText from a plain string — no page structure."""
    lines = [
        Line(text=ln.strip(), page=1)
        for ln in text.splitlines()
        if ln.strip()
    ]
    return ExtractedText(
        text=text,
        page_count=1,
        lines=lines,
        extraction_method=method,
    )
